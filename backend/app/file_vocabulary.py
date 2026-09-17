"""단어장 파일(PDF·CSV·XLSX)에서 영어 단어와 한국어 뜻을 뽑는다.

흐름: 업로드 검증·저장 → 텍스트 추출 → 줄 경계로 분할 → 조각마다 Codex → 병합.
Codex를 부르는 곳은 `extract_vocabulary_from_document` 하나라서 나머지는 Codex 없이 테스트한다.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fastapi import UploadFile

from .codex_cli import run_codex_exec


CHUNK_SIZE = 1024 * 1024
MAX_DOCUMENT_SIZE = 100 * 1024 * 1024
# 카드 200개 × 10장. 이보다 큰 목록은 파일을 나눠 올리라고 안내한다.
MAX_ENTRIES = 2000
# 조각 하나가 Codex 프롬프트 하나다. 30,000자면 CSV 기준 약 1,500행이라 한 번에 처리된다.
SEGMENT_CHARS = 30_000
# 이보다 짧으면 스캔 PDF(이미지만 있는 PDF)로 본다.
MIN_PDF_TEXT_CHARS = 200
MAX_WORD_LENGTH = 64
MAX_MEANING_LENGTH = 120

SUPPORTED_KINDS = ("pdf", "csv", "xlsx")
KIND_BY_EXTENSION = {".pdf": "pdf", ".csv": "csv", ".xlsx": "xlsx"}
UNSUPPORTED_FILE_DETAIL = "PDF, CSV, XLSX 파일만 올릴 수 있습니다."
LEGACY_EXCEL_DETAIL = "xls 파일은 읽을 수 없습니다. 엑셀에서 xlsx로 저장해서 올려 주세요."
TOO_LARGE_DETAIL = "100MB 이하 파일만 올릴 수 있습니다."
ENCRYPTED_PDF_DETAIL = "암호가 걸린 PDF는 읽을 수 없습니다."
SCANNED_PDF_DETAIL = "텍스트가 없는 PDF입니다. 스크린샷으로 '이미지에서 추가'를 써 주세요."
BROKEN_FILE_DETAIL = "파일을 읽지 못했습니다. 파일이 손상되지 않았는지 확인해 주세요."
NO_WORDS_DETAIL = "파일에서 영어 단어를 찾지 못했습니다."
INVALID_CODEX_JSON_DETAIL = "단어 목록을 정리하지 못했습니다. 잠시 뒤 다시 시도해 주세요."

ASCII_LETTER_PATTERN = re.compile(r"[A-Za-z]")


class DocumentError(Exception):
    """사용자에게 그대로 보여 줄 수 있는 실패."""

    def __init__(self, detail: str, *, status_code: int = 422) -> None:
        super().__init__(detail)
        self.detail = detail
        self.status_code = status_code


@dataclass(frozen=True)
class VocabularyEntry:
    word: str
    meaning_ko: str | None = None

    def as_dict(self) -> dict[str, str | None]:
        return {"word": self.word, "meaning_ko": self.meaning_ko}


@dataclass
class ExtractionResult:
    entries: list[VocabularyEntry] = field(default_factory=list)
    suggested_folder_name: str | None = None
    target_folder_name: str | None = None


# ---------------------------------------------------------------------------
# 업로드 검증·저장
# ---------------------------------------------------------------------------


def detect_document_kind(filename: str | None, header_bytes: bytes) -> str:
    """확장자와 파일 머리로 종류를 정한다. 못 정하면 DocumentError."""
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xls":
        raise DocumentError(LEGACY_EXCEL_DETAIL)

    kind = KIND_BY_EXTENSION.get(suffix)
    if kind is None:
        raise DocumentError(UNSUPPORTED_FILE_DETAIL)

    if kind == "pdf" and not header_bytes.startswith(b"%PDF"):
        raise DocumentError(UNSUPPORTED_FILE_DETAIL)
    if kind == "xlsx" and not header_bytes.startswith(b"PK\x03\x04"):
        raise DocumentError(UNSUPPORTED_FILE_DETAIL)

    return kind


async def save_validated_document_upload(
    file: UploadFile,
    destination: Path,
    *,
    max_size: int | None = None,
) -> str:
    """업로드를 검증하며 디스크에 쓴다. 종류('pdf'|'csv'|'xlsx')를 돌려준다."""
    max_size = max_size or MAX_DOCUMENT_SIZE
    first_chunk = await file.read(CHUNK_SIZE)
    try:
        kind = detect_document_kind(file.filename, first_chunk[:16])
    except DocumentError:
        await file.close()
        raise

    total_size = 0
    try:
        with destination.open("wb") as output_file:
            chunk = first_chunk
            while chunk:
                total_size += len(chunk)
                if total_size > max_size:
                    raise DocumentError(TOO_LARGE_DETAIL, status_code=413)
                output_file.write(chunk)
                chunk = await file.read(CHUNK_SIZE)
    except DocumentError:
        destination.unlink(missing_ok=True)
        raise
    finally:
        await file.close()

    if total_size == 0:
        destination.unlink(missing_ok=True)
        raise DocumentError(BROKEN_FILE_DETAIL)

    return kind


# ---------------------------------------------------------------------------
# 텍스트 추출
# ---------------------------------------------------------------------------


def _extract_pdf_text(path: Path) -> str:
    # pypdf 문서: PdfReader(stream), reader.is_encrypted, page.extract_text()
    # https://pypdf.readthedocs.io/en/stable/user/extract-text.html
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentError(ENCRYPTED_PDF_DETAIL)
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except Exception as exc:  # pypdf는 손상 파일에서 여러 예외 타입을 던진다
        raise DocumentError(BROKEN_FILE_DETAIL) from exc

    text = "\n".join(page.strip() for page in pages if page.strip())
    if len(text.strip()) < MIN_PDF_TEXT_CHARS:
        raise DocumentError(SCANNED_PDF_DETAIL)
    return text


def _decode_csv_bytes(raw: bytes) -> str:
    # 한국어 엑셀이 저장한 CSV는 cp949가 흔하다. BOM 있는 utf-8부터 시도한다.
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _rows_to_text(rows: Any) -> str:
    lines: list[str] = []
    for row in rows:
        cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
        if cells:
            lines.append("\t".join(cells))
    return "\n".join(lines)


def _extract_csv_text(path: Path) -> str:
    try:
        content = _decode_csv_bytes(path.read_bytes())
        return _rows_to_text(csv.reader(io.StringIO(content)))
    except (OSError, csv.Error) as exc:
        raise DocumentError(BROKEN_FILE_DETAIL) from exc


def _extract_xlsx_text(path: Path) -> str:
    # openpyxl 문서: load_workbook(read_only=True) 는 close() 를 직접 해야 한다.
    # https://openpyxl.readthedocs.io/en/stable/optimized.html
    # 경로를 주면 확장자(.xlsx)를 검사하는데 임시 파일에는 확장자가 없다. 파일 객체로 연다.
    from openpyxl import load_workbook

    try:
        with path.open("rb") as handle:
            workbook = load_workbook(handle, read_only=True, data_only=True)
            try:
                sheets = [_rows_to_text(sheet.iter_rows(values_only=True)) for sheet in workbook]
            finally:
                workbook.close()
    except OSError:
        raise
    except Exception as exc:
        raise DocumentError(BROKEN_FILE_DETAIL) from exc

    return "\n".join(sheet for sheet in sheets if sheet)


def extract_document_text(path: Path, kind: str) -> str:
    if kind == "pdf":
        return _extract_pdf_text(path)
    if kind == "csv":
        return _extract_csv_text(path)
    if kind == "xlsx":
        return _extract_xlsx_text(path)
    raise DocumentError(UNSUPPORTED_FILE_DETAIL)


# ---------------------------------------------------------------------------
# 분할·프롬프트·파싱
# ---------------------------------------------------------------------------


def split_text(text: str, max_chars: int | None = None) -> list[str]:
    """줄 경계에서 자른다. 한 줄이 max_chars 보다 길면 그 줄만 잘라 넣는다."""
    max_chars = max_chars or SEGMENT_CHARS
    segments: list[str] = []
    current: list[str] = []
    current_length = 0

    def flush() -> None:
        nonlocal current, current_length
        if current:
            segments.append("\n".join(current))
        current = []
        current_length = 0

    for line in text.splitlines():
        if not line.strip():
            continue
        while len(line) > max_chars:
            flush()
            segments.append(line[:max_chars])
            line = line[max_chars:]
        added = len(line) + (1 if current else 0)
        if current_length + added > max_chars:
            flush()
            added = len(line)
        current.append(line)
        current_length += added

    flush()
    return segments


def build_extraction_prompt(
    segment: str,
    *,
    file_name: str,
    segment_index: int,
    segment_count: int,
    user_message: str | None,
    folder_names: list[str],
) -> str:
    folder_list = json.dumps(folder_names, ensure_ascii=False)
    request = (user_message or "").strip() or "(no message)"
    return "\n".join(
        [
            "You are reading text extracted from a vocabulary-list file for VocaLoop, an English vocabulary app.",
            f'File name: "{file_name}". This is part {segment_index + 1} of {segment_count}.',
            "Do not inspect local files, run shell commands, or browse the web.",
            "Extract the English vocabulary words or short English phrases that this list teaches, in order of appearance.",
            "If a Korean meaning is written next to a word, copy it as a short gloss (1-3 terms). Otherwise use null.",
            "Ignore numbering, checkboxes, section headers, page numbers, dates, example sentences, and explanations.",
            "Remove duplicates case-insensitively. Do not guess unclear text. Do not invent meanings.",
            'Suggest a short folder name for this list (from its title, file name, or topic) in "suggested_folder_name".',
            f"The user's existing folders are: {folder_list}.",
            f'The user said: "{request}"',
            'If the user asks to put the words into one of those existing folders, set "target_folder_name" to that exact folder name. Otherwise use null.',
            "Return JSON only, no markdown fences, with this shape:",
            '{"entries":[{"word":"abate","meaning_ko":"줄이다"},{"word":"candid","meaning_ko":null}],'
            '"suggested_folder_name":"TOEFL Day 1","target_folder_name":null}',
            "",
            "TEXT:",
            segment,
        ]
    )


def _normalize_text(value: object, *, max_length: int) -> str | None:
    if value is None:
        return None
    normalized = re.sub(r"\s+", " ", str(value)).strip()
    if not normalized:
        return None
    return normalized[:max_length]


def normalize_entries(values: object, *, seen: set[str] | None = None, limit: int | None = None) -> list[VocabularyEntry]:
    """Codex가 준 entries를 정리한다. `seen`을 넘기면 조각 사이의 중복도 걸러진다."""
    if not isinstance(values, list):
        return []

    limit = limit or MAX_ENTRIES
    seen_keys = seen if seen is not None else set()
    entries: list[VocabularyEntry] = []
    for value in values:
        if isinstance(value, str):
            raw_word, raw_meaning = value, None
        elif isinstance(value, dict):
            raw_word, raw_meaning = value.get("word"), value.get("meaning_ko")
        else:
            continue

        word = _normalize_text(raw_word, max_length=MAX_WORD_LENGTH)
        if not word or not ASCII_LETTER_PATTERN.search(word):
            continue
        key = word.lower()
        if key in seen_keys:
            continue
        seen_keys.add(key)
        entries.append(VocabularyEntry(word=word, meaning_ko=_normalize_text(raw_meaning, max_length=MAX_MEANING_LENGTH)))
        if len(entries) >= limit:
            break

    return entries


def parse_extraction_output(text: str) -> dict[str, Any]:
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentError(INVALID_CODEX_JSON_DETAIL, status_code=502) from exc
    if not isinstance(parsed, dict):
        raise DocumentError(INVALID_CODEX_JSON_DETAIL, status_code=502)
    return parsed


def resolve_folder_name(candidate: object, folder_names: list[str]) -> str | None:
    """Codex가 돌려준 이름을 기존 폴더명과 대소문자·공백 무시로 맞춘다. 없으면 None."""
    normalized = _normalize_text(candidate, max_length=255)
    if not normalized:
        return None
    wanted = normalized.lower()
    for name in folder_names:
        if name.strip().lower() == wanted:
            return name
    return None


# ---------------------------------------------------------------------------
# 전체 흐름
# ---------------------------------------------------------------------------


def extract_vocabulary_from_document(
    path: Path,
    kind: str,
    *,
    file_name: str,
    user_message: str | None,
    folder_names: list[str],
    run: Callable[[str], str] = run_codex_exec,
) -> ExtractionResult:
    text = extract_document_text(path, kind)
    segments = split_text(text)
    if not segments:
        raise DocumentError(NO_WORDS_DETAIL)

    result = ExtractionResult()
    seen: set[str] = set()
    for index, segment in enumerate(segments):
        prompt = build_extraction_prompt(
            segment,
            file_name=file_name,
            segment_index=index,
            segment_count=len(segments),
            user_message=user_message,
            folder_names=folder_names,
        )
        parsed = parse_extraction_output(run(prompt))
        remaining = MAX_ENTRIES - len(result.entries)
        result.entries.extend(normalize_entries(parsed.get("entries"), seen=seen, limit=remaining))
        if result.suggested_folder_name is None:
            result.suggested_folder_name = _normalize_text(parsed.get("suggested_folder_name"), max_length=30)
        if result.target_folder_name is None:
            result.target_folder_name = resolve_folder_name(parsed.get("target_folder_name"), folder_names)
        if len(result.entries) >= MAX_ENTRIES:
            break

    if not result.entries:
        raise DocumentError(NO_WORDS_DETAIL)

    return result
